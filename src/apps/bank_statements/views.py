from django.shortcuts import render
from django.http import HttpResponse
from rest_framework import viewsets, status, filters
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django_filters.rest_framework import DjangoFilterBackend
from drf_yasg.utils import swagger_auto_schema
from drf_yasg import openapi
import logging
import traceback
import csv
import io
from decimal import Decimal
from datetime import datetime

from apps.accounts.models import BusinessAccount
from .models import BankStatement, BankTransaction, TransactionCategory, MOROCCAN_BANKS
from .serializers import (
    BankStatementSerializer,
    BankStatementDetailSerializer,
    BankTransactionSerializer,
    TransactionCategorySerializer
)
from .statement_parser import StatementParser
from .utils import TransactionMapper

logger = logging.getLogger(__name__)


class BankStatementViewSet(viewsets.ModelViewSet):
    """
    ViewSet for managing bank statements
    """
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['account_holder', 'bank_name', 'account_number', 'branch']
    ordering_fields = ['period_start', 'period_end', 'created_at', 'opening_balance', 'closing_balance']
    filterset_fields = ['status', 'currency', 'parser_used']

    def get_serializer_class(self):
        """Return different serializer for list and detail"""
        if self.action == 'retrieve':
            return BankStatementDetailSerializer
        return BankStatementSerializer

    def get_queryset(self):
        """Return bank statements for the user's business account"""
        if not self.request.user.is_authenticated:
            return BankStatement.objects.none()
        
        # Get the business accounts the user has access to
        business_accounts = self.request.user.business_accounts.all()
        if hasattr(self.request.user, 'owned_business'):
            business_accounts = business_accounts | BusinessAccount.objects.filter(id=self.request.user.owned_business.id)
        
        return BankStatement.objects.filter(business_account__in=business_accounts)
    
    def perform_create(self, serializer):
        """Set the business account when creating a bank statement"""
        if hasattr(self.request.user, 'owned_business'):
            serializer.save(business_account=self.request.user.owned_business)
        else:
            return Response(
                {"error": "User does not have a business account"},
                status=status.HTTP_400_BAD_REQUEST
            )

    @swagger_auto_schema(
        operation_summary="Get list of Moroccan banks",
        operation_description="Returns the list of available Moroccan banks for selection",
        responses={200: openapi.Response(description="List of banks")},
    )
    @action(detail=False, methods=['get'])
    def moroccan_banks(self, request):
        """
        Get list of Moroccan banks
        """
        banks = [{'code': code, 'name': name} for code, name in MOROCCAN_BANKS]
        return Response(banks, status=status.HTTP_200_OK)

    @swagger_auto_schema(
        operation_summary="Parse bank statement file",
        operation_description="Extract transactions and metadata from uploaded bank statement file (PDF, Image, CSV, Excel)",
        request_body=openapi.Schema(
            type=openapi.TYPE_OBJECT,
            required=['file'],
            properties={
                'file': openapi.Schema(type=openapi.TYPE_FILE, description='Bank statement file'),
                'use_gpt': openapi.Schema(type=openapi.TYPE_BOOLEAN, description='Use GPT parser (default: true)', default=True),
            }
        ),
        responses={
            200: openapi.Response(description="Extracted bank statement data"),
            400: openapi.Response(description="Bad request"),
            500: openapi.Response(description="Server error"),
        },
    )
    @action(detail=False, methods=['post'])
    def parse_statement(self, request):
        """
        Extract data from uploaded bank statement file
        """
        try:
            if 'file' not in request.FILES:
                return Response({"error": "No file provided"}, status=status.HTTP_400_BAD_REQUEST)
            
            statement_file = request.FILES['file']
            use_gpt = request.data.get('use_gpt', 'true').lower() in ('true', 't', 'yes', 'y', '1')
            
            # Parse the statement
            # Note: Poppler and Tesseract are only needed for image files and scanned PDFs
            # Most PDFs with text will work fine without these dependencies
            try:
                extracted_data = StatementParser.parse_statement(statement_file, use_gpt=use_gpt)
                return Response(extracted_data, status=status.HTTP_200_OK)
            except ValueError as e:
                # Check if it's a dependency issue
                error_msg = str(e)
                if "Poppler" in error_msg or "Tesseract" in error_msg:
                    instructions = StatementParser.get_dependency_installation_instructions()
                    return Response({
                        "error": error_msg,
                        "installation_instructions": instructions,
                        "note": "Text-based PDFs can be processed without these dependencies. Only image files and scanned PDFs require Poppler and Tesseract."
                    }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
                return Response(
                    {"error": error_msg},
                    status=status.HTTP_400_BAD_REQUEST
                )
                
        except Exception as e:
            logger.error(f"Bank statement parsing error: {str(e)}\n{traceback.format_exc()}")
            return Response(
                {"error": f"Failed to parse bank statement: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @swagger_auto_schema(
        operation_summary="Create statement from parsed data",
        operation_description="Create a bank statement and transactions from previously parsed data",
        request_body=openapi.Schema(
            type=openapi.TYPE_OBJECT,
            required=['statement_data'],
            properties={
                'statement_data': openapi.Schema(
                    type=openapi.TYPE_OBJECT,
                    description='Parsed statement data from parse_statement endpoint'
                ),
                'file': openapi.Schema(type=openapi.TYPE_FILE, description='Original statement file (optional)'),
            }
        ),
        responses={
            201: BankStatementDetailSerializer(),
            400: openapi.Response(description="Bad request"),
        },
    )
    @action(detail=False, methods=['post'])
    def create_from_parsed(self, request):
        """
        Create a bank statement and transactions from parsed data
        """
        try:
            statement_data = request.data.get('statement_data')
            if not statement_data:
                return Response(
                    {"error": "statement_data is required"},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # Parse JSON string if needed
            if isinstance(statement_data, str):
                import json
                statement_data = json.loads(statement_data)
            
            if not hasattr(request.user, 'owned_business'):
                return Response(
                    {"error": "User does not have a business account"},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # The data is already flat (not nested under 'statement' key)
            # because the frontend extracts it from the parse response
            stmt = statement_data
            
            # Initialize transaction mapper for standardized data handling
            mapper = TransactionMapper()
            
            # Determine file type
            file_type = None
            if 'file' in request.FILES:
                file_name = request.FILES['file'].name.lower()
                if file_name.endswith('.pdf'):
                    file_type = 'PDF'
                elif file_name.endswith(('.jpg', '.jpeg', '.png')):
                    file_type = 'Image'
                elif file_name.endswith('.csv'):
                    file_type = 'CSV'
                elif file_name.endswith(('.xls', '.xlsx')):
                    file_type = 'Excel'
            
            # Create bank statement
            bank_statement = BankStatement.objects.create(
                business_account=request.user.owned_business,
                account_holder=stmt.get('account_holder'),
                bank_name=stmt.get('bank_name') or 'Auto-detected',
                account_number=stmt.get('account_number'),
                branch=stmt.get('branch'),
                period_start=stmt.get('period_start'),
                period_end=stmt.get('period_end'),
                opening_balance=Decimal(str(stmt.get('opening_balance', 0))) if stmt.get('opening_balance') else None,
                closing_balance=Decimal(str(stmt.get('closing_balance', 0))) if stmt.get('closing_balance') else None,
                currency=stmt.get('currency') or 'MAD',
                file_type=file_type,
                status='processing',
                parser_used='gpt'
            )
            
            # Save original file if provided
            if 'file' in request.FILES:
                bank_statement.original_file = request.FILES['file']
                bank_statement.save()
            
            # Create transactions using the mapper for standardization
            raw_transactions = stmt.get('transactions', [])
            created_transactions = []
            
            for raw_tx in raw_transactions:
                try:
                    # Normalize transaction data using the mapper
                    normalized = mapper.normalize_transaction(raw_tx)
                    
                    transaction = BankTransaction.objects.create(
                        bank_statement=bank_statement,
                        operation_date=normalized['operation_date'],
                        value_date=normalized['value_date'],
                        reference=normalized['reference'],
                        description=normalized['description'],
                        debit=normalized['debit'],
                        credit=normalized['credit'],
                        category=normalized['category'],
                        merchant=normalized['merchant'],
                        location=normalized['location'],
                        beneficiary=normalized['beneficiary'],
                        depositor=normalized['depositor'],
                        counterparty_name=normalized['counterparty_name'],
                        counterparty_account=normalized['counterparty_account'],
                        time=normalized['time'],
                        payment_method=normalized['payment_method'],
                    )
                    created_transactions.append(transaction)
                except Exception as e:
                    logger.warning(f"Failed to create transaction: {str(e)}")
                    continue
            
            # Calculate totals
            bank_statement.calculate_totals()
            bank_statement.status = 'completed'
            bank_statement.save()
            
            # Return created statement
            serializer = BankStatementDetailSerializer(bank_statement)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
            
        except Exception as e:
            logger.error(f"Error creating statement from parsed data: {str(e)}\n{traceback.format_exc()}")
            return Response(
                {"error": f"Failed to create statement: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @swagger_auto_schema(
        operation_summary="Recalculate statement totals",
        operation_description="Recalculate total debits, credits, and transaction count",
        responses={200: openapi.Response(description="Totals recalculated")},
    )
    @action(detail=True, methods=['post'])
    def recalculate_totals(self, request, pk=None):
        """Recalculate statement totals"""
        statement = self.get_object()
        statement.calculate_totals()
        return Response({"status": "Totals recalculated successfully"}, status=status.HTTP_200_OK)


class BankTransactionViewSet(viewsets.ModelViewSet):
    """
    ViewSet for managing bank transactions
    """
    serializer_class = BankTransactionSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['description', 'reference', 'merchant', 'beneficiary', 'depositor']
    ordering_fields = ['operation_date', 'value_date', 'debit', 'credit', 'created_at']
    filterset_fields = ['category', 'is_reconciled', 'needs_review', 'is_duplicate']

    def get_queryset(self):
        """Return transactions for statements the user has access to"""
        if not self.request.user.is_authenticated:
            return BankTransaction.objects.none()
        
        # Get the business accounts the user has access to
        business_accounts = self.request.user.business_accounts.all()
        if hasattr(self.request.user, 'owned_business'):
            business_accounts = business_accounts | BusinessAccount.objects.filter(id=self.request.user.owned_business.id)
        
        # Get statements for these business accounts
        statements = BankStatement.objects.filter(business_account__in=business_accounts)
        
        # Filter by statement_id if provided
        statement_id = self.request.query_params.get('statement_id')
        if statement_id:
            statements = statements.filter(id=statement_id)
        
        return BankTransaction.objects.filter(bank_statement__in=statements)

    @swagger_auto_schema(
        operation_summary="Mark transaction as reconciled",
        operation_description="Mark a transaction as reconciled",
        responses={200: openapi.Response(description="Transaction reconciled")},
    )
    @action(detail=True, methods=['post'])
    def reconcile(self, request, pk=None):
        """Mark transaction as reconciled"""
        transaction = self.get_object()
        transaction.is_reconciled = True
        transaction.save()
        return Response({"status": "Transaction marked as reconciled"}, status=status.HTTP_200_OK)

    @swagger_auto_schema(
        operation_summary="Mark transaction as duplicate",
        operation_description="Mark a transaction as a duplicate",
        responses={200: openapi.Response(description="Transaction marked as duplicate")},
    )
    @action(detail=True, methods=['post'])
    def mark_duplicate(self, request, pk=None):
        """Mark transaction as duplicate"""
        transaction = self.get_object()
        transaction.is_duplicate = True
        transaction.save()
        return Response({"status": "Transaction marked as duplicate"}, status=status.HTTP_200_OK)

    @swagger_auto_schema(
        operation_summary="Get suggested invoice matches for a transaction",
        operation_description="Find invoices that potentially match this bank transaction based on amount, date, and description analysis",
        responses={
            200: openapi.Response(
                description="List of suggested invoice matches",
                schema=openapi.Schema(
                    type=openapi.TYPE_OBJECT,
                    properties={
                        'matches': openapi.Schema(
                            type=openapi.TYPE_ARRAY,
                            items=openapi.Schema(
                                type=openapi.TYPE_OBJECT,
                                properties={
                                    'invoice_id': openapi.Schema(type=openapi.TYPE_STRING),
                                    'invoice_number': openapi.Schema(type=openapi.TYPE_STRING),
                                    'invoice_type': openapi.Schema(type=openapi.TYPE_STRING, enum=['customer', 'supplier']),
                                    'total': openapi.Schema(type=openapi.TYPE_NUMBER),
                                    'due_date': openapi.Schema(type=openapi.TYPE_STRING, format='date'),
                                    'party_name': openapi.Schema(type=openapi.TYPE_STRING),
                                    'score': openapi.Schema(type=openapi.TYPE_INTEGER, description='Match score 0-100'),
                                    'confidence_level': openapi.Schema(type=openapi.TYPE_STRING, enum=['high', 'medium', 'low']),
                                    'reasons': openapi.Schema(type=openapi.TYPE_ARRAY, items=openapi.Schema(type=openapi.TYPE_STRING)),
                                }
                            )
                        ),
                        'transaction': openapi.Schema(type=openapi.TYPE_OBJECT, description='Transaction details'),
                    }
                )
            ),
            404: openapi.Response(description="Transaction not found"),
        },
    )
    @action(detail=True, methods=['get'])
    def suggested_matches(self, request, pk=None):
        """
        Get suggested invoice matches for a transaction.
        Uses intelligent matching based on:
        - Amount matching (exact or close)
        - Date proximity
        - Reference/invoice number matching
        - Counterparty name matching
        """
        from .services import InvoiceMatcher, MATCH_REASON_LABELS
        
        transaction = self.get_object()
        
        # Get the business account for invoice lookup
        business_account = transaction.bank_statement.business_account
        
        # Initialize matcher and find matches
        matcher = InvoiceMatcher(business_account)
        raw_matches = matcher.find_matches(transaction)
        
        # Format matches for JSON response
        formatted_matches = []
        for match in raw_matches:
            invoice = match['invoice']
            # Determine party name based on invoice type
            if invoice.invoice_type == 'purchase':
                party_name = invoice.supplier.corporate_name if invoice.supplier else 'Unknown Supplier'
                invoice_type_display = 'supplier'
            else:
                party_name = invoice.customer.corporate_name if invoice.customer else 'Unknown Customer'
                invoice_type_display = 'customer'
            
            formatted_matches.append({
                'invoice_id': str(invoice.id),
                'invoice_number': invoice.invoice_number or '-',
                'invoice_type': invoice_type_display,
                'total': float(invoice.total) if invoice.total else 0,
                'due_date': invoice.due_date.isoformat() if invoice.due_date else None,
                'party_name': party_name,
                'score': match['score'],
                'confidence_level': match['confidence'],
                'reasons': [MATCH_REASON_LABELS.get(r, r) for r in match['reasons']],
            })
        
        # Format transaction info for response
        transaction_info = {
            'id': str(transaction.id),
            'operation_date': transaction.operation_date.isoformat() if transaction.operation_date else None,
            'description': transaction.description,
            'debit': float(transaction.debit) if transaction.debit else None,
            'credit': float(transaction.credit) if transaction.credit else None,
            'amount': float(transaction.debit or transaction.credit or 0),
            'is_debit': transaction.debit is not None and transaction.debit > 0,
            'counterparty_name': transaction.counterparty_name,
            'reference': transaction.reference,
            'is_reconciled': transaction.is_reconciled,
            'matched_invoice_id': str(transaction.matched_invoice.id) if transaction.matched_invoice else None,
        }
        
        return Response({
            'transaction': transaction_info,
            'matches': formatted_matches,
        }, status=status.HTTP_200_OK)

    @swagger_auto_schema(
        operation_summary="Reconcile transaction with an invoice",
        operation_description="Link a bank transaction to an invoice and mark the invoice as paid",
        request_body=openapi.Schema(
            type=openapi.TYPE_OBJECT,
            required=['invoice_id'],
            properties={
                'invoice_id': openapi.Schema(type=openapi.TYPE_STRING, description='UUID of the invoice to link'),
                'notes': openapi.Schema(type=openapi.TYPE_STRING, description='Optional notes about the reconciliation'),
            }
        ),
        responses={
            200: openapi.Response(
                description="Transaction reconciled successfully",
                schema=openapi.Schema(
                    type=openapi.TYPE_OBJECT,
                    properties={
                        'status': openapi.Schema(type=openapi.TYPE_STRING),
                        'transaction_id': openapi.Schema(type=openapi.TYPE_STRING),
                        'invoice_id': openapi.Schema(type=openapi.TYPE_STRING),
                        'invoice_status': openapi.Schema(type=openapi.TYPE_STRING),
                    }
                )
            ),
            400: openapi.Response(description="Bad request"),
            404: openapi.Response(description="Transaction or invoice not found"),
        },
    )
    @action(detail=True, methods=['post'])
    def reconcile_with_invoice(self, request, pk=None):
        """
        Reconcile a transaction with an invoice.
        This will:
        1. Link the transaction to the invoice
        2. Mark the transaction as reconciled
        3. Update the invoice status to 'paid'
        4. Set the invoice paid_date
        """
        from .services import InvoiceMatcher
        from apps.invoices.models import Invoice
        
        transaction = self.get_object()
        invoice_id = request.data.get('invoice_id')
        notes = request.data.get('notes', '')
        
        if not invoice_id:
            return Response(
                {"error": "invoice_id is required"},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        # Get the invoice
        try:
            invoice = Invoice.objects.get(id=invoice_id)
        except Invoice.DoesNotExist:
            return Response(
                {"error": "Invoice not found"},
                status=status.HTTP_404_NOT_FOUND
            )
        
        # Verify the invoice belongs to the same business account
        business_account = transaction.bank_statement.business_account
        if invoice.business_account != business_account:
            return Response(
                {"error": "Invoice does not belong to the same business account"},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        # Check if already reconciled with a different invoice
        if transaction.matched_invoice and transaction.matched_invoice.id != invoice.id:
            return Response(
                {"error": "Transaction is already reconciled with a different invoice"},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        # Perform reconciliation using the matcher service
        matcher = InvoiceMatcher(business_account)
        try:
            result = matcher.reconcile_transaction_with_invoice(
                transaction=transaction,
                invoice=invoice,
                user=request.user
            )
            
            return Response({
                'status': 'Transaction reconciled successfully',
                'transaction_id': str(transaction.id),
                'invoice_id': str(invoice.id),
                'invoice_number': invoice.invoice_number,
                'invoice_status': invoice.status,
                'reconciled_at': result.get('reconciled_at'),
            }, status=status.HTTP_200_OK)
            
        except Exception as e:
            logger.error(f"Reconciliation error: {str(e)}\n{traceback.format_exc()}")
            return Response(
                {"error": f"Failed to reconcile: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @swagger_auto_schema(
        operation_summary="Unreconcile a transaction",
        operation_description="Remove the link between a transaction and its matched invoice",
        responses={
            200: openapi.Response(description="Transaction unreconciled successfully"),
            400: openapi.Response(description="Transaction is not reconciled"),
        },
    )
    @action(detail=True, methods=['post'])
    def unreconcile(self, request, pk=None):
        """
        Remove reconciliation from a transaction.
        This will also reset the invoice status back to unpaid.
        """
        from .services import InvoiceMatcher
        
        transaction = self.get_object()
        
        if not transaction.is_reconciled and not transaction.matched_invoice:
            return Response(
                {"error": "Transaction is not reconciled"},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        # Store previous invoice info for response
        previous_invoice_id = str(transaction.matched_invoice.id) if transaction.matched_invoice else None
        previous_invoice_number = transaction.matched_invoice.invoice_number if transaction.matched_invoice else None
        
        # Use the InvoiceMatcher service to properly unreconcile
        business_account = transaction.bank_statement.business_account
        matcher = InvoiceMatcher(business_account)
        
        try:
            result = matcher.unreconcile_transaction(transaction, request.user)
            
            return Response({
                'status': 'Transaction unreconciled successfully',
                'transaction_id': str(transaction.id),
                'previous_invoice_id': previous_invoice_id,
                'previous_invoice_number': previous_invoice_number,
                'invoice_status_reset': True,
            }, status=status.HTTP_200_OK)
        except Exception as e:
            logger.error(f"Unreconcile error: {str(e)}\n{traceback.format_exc()}")
            return Response(
                {"error": f"Failed to unreconcile: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @swagger_auto_schema(
        operation_summary="Export bank transactions",
        operation_description="Export bank transactions to CSV or Excel format for accountants",
        manual_parameters=[
            openapi.Parameter(
                'format', openapi.IN_QUERY,
                description="Export format: csv or xlsx",
                type=openapi.TYPE_STRING,
                enum=['csv', 'xlsx'],
                default='csv'
            ),
            openapi.Parameter(
                'start_date', openapi.IN_QUERY,
                description="Filter transactions from this date (YYYY-MM-DD)",
                type=openapi.TYPE_STRING,
                format='date'
            ),
            openapi.Parameter(
                'end_date', openapi.IN_QUERY,
                description="Filter transactions up to this date (YYYY-MM-DD)",
                type=openapi.TYPE_STRING,
                format='date'
            ),
            openapi.Parameter(
                'statement_id', openapi.IN_QUERY,
                description="Filter by specific bank statement",
                type=openapi.TYPE_STRING
            ),
            openapi.Parameter(
                'reconciled_only', openapi.IN_QUERY,
                description="Export only reconciled transactions (linked to invoices)",
                type=openapi.TYPE_BOOLEAN,
                default=False
            ),
        ],
        responses={
            200: openapi.Response(description="File download"),
            400: openapi.Response(description="Bad request"),
        },
    )
    @action(detail=False, methods=['get'], url_path='export', url_name='export')
    def export(self, request):
        """
        Export bank transactions to CSV or Excel format.
        Columns: Date Opération, Date Valeur, Description, Référence, Contrepartie,
                 Débit, Crédit, Solde, Catégorie, Rapproché, N° Facture
        """
        # Get parameters
        export_format = request.query_params.get('format', 'csv').lower()
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')
        statement_id = request.query_params.get('statement_id')
        reconciled_only = request.query_params.get('reconciled_only', 'false').lower() in ('true', '1', 'yes')
        
        # Get base queryset
        queryset = self.get_queryset()
        
        # Apply filters
        if start_date:
            try:
                start = datetime.strptime(start_date, '%Y-%m-%d').date()
                queryset = queryset.filter(operation_date__gte=start)
            except ValueError:
                return Response(
                    {"error": "Invalid start_date format. Use YYYY-MM-DD"},
                    status=status.HTTP_400_BAD_REQUEST
                )
        
        if end_date:
            try:
                end = datetime.strptime(end_date, '%Y-%m-%d').date()
                queryset = queryset.filter(operation_date__lte=end)
            except ValueError:
                return Response(
                    {"error": "Invalid end_date format. Use YYYY-MM-DD"},
                    status=status.HTTP_400_BAD_REQUEST
                )
        
        if statement_id:
            queryset = queryset.filter(bank_statement_id=statement_id)
        
        if reconciled_only:
            queryset = queryset.filter(is_reconciled=True, matched_invoice__isnull=False)
        
        # Order by operation date
        queryset = queryset.order_by('operation_date', 'created_at')
        
        # Prepare data
        transactions = list(queryset.select_related('matched_invoice'))
        
        # Column headers (French for accountants)
        headers = [
            'Date Opération',
            'Date Valeur',
            'Description',
            'Référence',
            'Contrepartie',
            'Débit',
            'Crédit',
            'Solde',
            'Catégorie',
            'Rapproché',
            'N° Facture'
        ]
        
        # Format date for export (DD/MM/YYYY - French format)
        def format_date(date_val):
            if date_val:
                return date_val.strftime('%d/%m/%Y')
            return ''
        
        # Format decimal for export
        def format_decimal(val):
            if val is not None and val > 0:
                return str(val).replace('.', ',')  # French decimal separator
            return ''
        
        # Build rows
        rows = []
        for tx in transactions:
            row = [
                format_date(tx.operation_date),
                format_date(tx.value_date),
                tx.description or '',
                tx.reference or '',
                tx.counterparty_name or '',
                format_decimal(tx.debit),
                format_decimal(tx.credit),
                format_decimal(tx.balance_after) if tx.balance_after else '',
                tx.category or '',
                'Oui' if tx.is_reconciled else 'Non',
                tx.matched_invoice.invoice_number if tx.matched_invoice else '',
            ]
            rows.append(row)
        
        # Generate filename
        business_name = 'export'
        if hasattr(request.user, 'owned_business') and request.user.owned_business:
            business_name = request.user.owned_business.name or 'export'
        business_name = business_name.replace(' ', '_')
        
        date_suffix = datetime.now().strftime('%Y%m%d')
        
        if export_format == 'xlsx':
            # Excel export
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
                from openpyxl.utils import get_column_letter
            except ImportError:
                return Response(
                    {"error": "openpyxl is required for Excel export. Install with: pip install openpyxl"},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Mouvements Bancaires"
            
            # Header styling
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
            
            # Write headers
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border
            
            # Write data
            for row_idx, row in enumerate(rows, 2):
                for col_idx, value in enumerate(row, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = thin_border
            
            # Auto-adjust column widths
            for col_idx, header in enumerate(headers, 1):
                max_length = len(header)
                for row in rows:
                    cell_value = str(row[col_idx - 1]) if col_idx - 1 < len(row) else ''
                    max_length = max(max_length, len(cell_value))
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)
            
            # Save to buffer
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            
            filename = f"mouvements_bancaires_{business_name}_{date_suffix}.xlsx"
            response = HttpResponse(
                output.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response
        
        else:
            # CSV export (default)
            output = io.StringIO()
            writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_MINIMAL)
            
            # Write BOM for Excel UTF-8 compatibility
            output.write('\ufeff')
            
            # Write headers
            writer.writerow(headers)
            
            # Write data
            for row in rows:
                writer.writerow(row)
            
            filename = f"mouvements_bancaires_{business_name}_{date_suffix}.csv"
            response = HttpResponse(
                output.getvalue(),
                content_type='text/csv; charset=utf-8'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response


class TransactionCategoryViewSet(viewsets.ModelViewSet):
    """
    ViewSet for managing transaction categories
    """
    serializer_class = TransactionCategorySerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['name', 'description']
    ordering_fields = ['name', 'created_at']

    def get_queryset(self):
        """Return categories for the user's business account"""
        if not self.request.user.is_authenticated:
            return TransactionCategory.objects.none()
        
        # Get the business accounts the user has access to
        business_accounts = self.request.user.business_accounts.all()
        if hasattr(self.request.user, 'owned_business'):
            business_accounts = business_accounts | BusinessAccount.objects.filter(id=self.request.user.owned_business.id)
        
        return TransactionCategory.objects.filter(business_account__in=business_accounts)
    
    def perform_create(self, serializer):
        """Set the business account when creating a category"""
        if hasattr(self.request.user, 'owned_business'):
            serializer.save(business_account=self.request.user.owned_business)
        else:
            return Response(
                {"error": "User does not have a business account"},
                status=status.HTTP_400_BAD_REQUEST
            )
